import io
import pdfplumber
import docx


SUPPORTED_MIME_TYPES = {
    # Google Workspace (exported as plain text by drive_client)
    "application/vnd.google-apps.document",
    "application/vnd.google-apps.spreadsheet",
    "application/vnd.google-apps.presentation",
    # Standard office formats
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    # Plain text / CSV
    "text/plain",
    "text/csv",
}


def extract_text(file_bytes: bytes, mime_type: str, file_name: str = "") -> str:
    try:
        if mime_type == "application/pdf":
            return _extract_pdf(file_bytes)
        if mime_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ):
            return _extract_docx(file_bytes)
        if mime_type in (
            "application/vnd.google-apps.document",
            "application/vnd.google-apps.spreadsheet",
            "application/vnd.google-apps.presentation",
            "text/plain",
            "text/csv",
        ):
            return file_bytes.decode("utf-8", errors="replace")
        # Fallback: try decoding as UTF-8
        return file_bytes.decode("utf-8", errors="replace")
    except Exception as exc:
        return f"[EXTRACTION ERROR for '{file_name}': {exc}]"


def _extract_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    # Also pull text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)
